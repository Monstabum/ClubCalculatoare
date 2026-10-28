# -*- coding: utf-8 -*-
"""Generează CAIETUL DE SARCINI (Word) pentru aplicația AppClubCalculatoare,
în structura modelului instituțional (caiet de sarcini semnat) și în
formatarea cerută de modelul de raport (Times New Roman 12, 1.5, margini A4).

Autori: Dudnic Alexandru, Marinescu Mihail, Scripnic Vladlen.
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"D:\Club de Calculatoare"
IMG = os.path.join(BASE, "Testare_Dovezi")
OUT = os.path.join(BASE, "Caiet_de_Sarcini_AppClubCalculatoare.docx")

doc = Document()

# ------------------------------------------------------------------
#  Aspect general (conform modelului de raport: TNR 12, 1.5, A4)
# ------------------------------------------------------------------
for s in doc.sections:
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(2.0)
    s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(3.0)
    s.right_margin = Cm(1.0)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
rpr = normal.element.get_or_add_rPr()
rfonts = rpr.find(qn("w:rFonts"))
if rfonts is None:
    rfonts = OxmlElement("w:rFonts")
    rpr.append(rfonts)
for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rfonts.set(qn(attr), "Times New Roman")
lang = OxmlElement("w:lang")
lang.set(qn("w:val"), "ro-RO")
rpr.append(lang)

# titluri
for name, size in (("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12),
                   ("Heading 4", 12)):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string("000000")
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Bullet 2"):
    st = doc.styles[name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(2)


# ------------------------------------------------------------------
#  Funcții ajutătoare
# ------------------------------------------------------------------
def p(text, bold=False, italic=False, size=None, align=None, color=None,
      space_after=6, space_before=0, indent=None):
    par = doc.add_paragraph()
    r = par.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = "Times New Roman"
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if align == "center":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "right":
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if indent is not None:
        par.paragraph_format.left_indent = Cm(indent)
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.space_before = Pt(space_before)
    return par


def j(text, **kw):
    kw.setdefault("align", "justify")
    return p(text, **kw)


def bullet(text, level=1):
    style = "List Bullet" if level == 1 else "List Bullet 2"
    par = doc.add_paragraph(style=style)
    r = par.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    par.paragraph_format.line_spacing = 1.5
    par.paragraph_format.space_after = Pt(2)
    return par


def h(text, level=1):
    return doc.add_heading(text, level=level)


def table(rows, widths=None, size=11, header=True, align_center=False):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for jj, val in enumerate(row):
            cells[jj].text = ""
            para = cells[jj].paragraphs[0]
            para.paragraph_format.line_spacing = 1.0
            para.paragraph_format.space_after = Pt(1)
            run = para.add_run(str(val))
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            if header and i == 0:
                run.bold = True
                if align_center:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if header and i == 0:
            for c in t.rows[0].cells:
                sh = OxmlElement("w:shd")
                sh.set(qn("w:fill"), "D9D9D9")
                c._tc.get_or_add_tcPr().append(sh)
    if widths:
        for row in t.rows:
            for idx, w in enumerate(widths):
                row.cells[idx].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


CROP_DIR = os.path.join(BASE, "Testare_Dovezi", "_crop_caiet")


def _crop_left(src, left=158):
    """Decupează banda de desktop din stânga capturii, păstrând fereastra aplicației."""
    try:
        from PIL import Image
    except Exception:
        return src
    try:
        os.makedirs(CROP_DIR, exist_ok=True)
        im = Image.open(src)
        w, h = im.size
        if left <= 0 or left >= w - 200:
            return src
        out = os.path.join(CROP_DIR, os.path.basename(src))
        im.crop((left, 0, w, h)).save(out)
        return out
    except Exception:
        return src


def figure(path, caption, width_cm=15.5):
    full = os.path.join(IMG, path)
    if os.path.exists(full):
        full = _crop_left(full)
        doc.add_picture(full, width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.line_spacing = 1.0


def add_toc():
    par = doc.add_paragraph()
    run = par.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve")
    i.text = 'TOC \\o "1-3" \\h \\z \\u'
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    tt = OxmlElement("w:t")
    tt.text = "Cuprinsul se actualizează cu click dreapta > Update Field (sau F9)."
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, i, s, tt, e):
        run._r.append(el)


def enable_update_fields():
    st = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    st.append(uf)


def add_footer_page():
    footer = doc.sections[0].footer
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = "PAGE"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, i, e):
        run._r.append(el)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)


# ==================================================================
#  PAGINA DE TITLU
# ==================================================================
p("Ministerul Educației și Cercetării al Republicii Moldova",
  size=12, align="center", space_after=2)
p("IP Colegiul Tehnic Feroviar din Bălți", size=12, align="center", space_after=2)
for _ in range(6):
    doc.add_paragraph()
p("CAIET DE SARCINI", bold=True, size=22, align="center", space_after=6)
p("Elaborarea sistemului informațional de gestionare a clubului de calculatoare "
  "„AppClubCalculatoare”", bold=True, size=14, align="center", space_after=24)
for _ in range(4):
    doc.add_paragraph()
tbl = table([
    ["Specialitatea:", "061210 Administrarea aplicațiilor web"],
    ["Calificarea:", "Tehnician de site-uri web"],
    ["Elaborat:", "Dudnic Alexandru, Marinescu Mihail, Scripnic Vladlen"],
    ["Grupa:", "__________________"],
    ["Verificat:", "Vasilos Corina"],
    ["Nota:", "__________________"],
], widths=[4.0, 11.0], header=False)
for _ in range(2):
    doc.add_paragraph()
p("Bălți, 2026", size=12, align="center", space_after=0)
doc.add_page_break()

# ==================================================================
#  CUPRINS
# ==================================================================
p("CUPRINS", bold=True, size=14, align="center", space_after=10)
add_toc()
doc.add_page_break()

# ==================================================================
#  1. INTRODUCERE
# ==================================================================
h("1. INTRODUCERE", 1)
j("Soluția software „Sistemul informațional de gestionare a clubului de calculatoare "
  "AppClubCalculatoare” (în continuare – SIA CCA sau „Aplicația”) va fi o soluție "
  "informatică de tip desktop destinată evidenței și gestionării activității unui club de "
  "calculatoare (sală de gaming / internet café): săli, operatori, inventarul de echipamente, "
  "membri, sesiuni de utilizare și rezervări, împreună cu încasările, rapoartele și controlul "
  "accesului utilizatorilor.")
j("Prezentul document definește cerințele privind funcționalitățile sistemului, arhitectura "
  "tehnologică, procesele operaționale, modulele componente, cerințele nefuncționale, cerințele "
  "față de echipa de implementare, etapele de desfășurare a lucrărilor, precum și livrabilele "
  "proiectului. Caietul de sarcini este documentul de referință pentru proiectare, dezvoltare, "
  "testare și acceptanță.")

h("1.1. Descrierea necesității", 2)
j("Necesitatea prezentei achiziții derivă din activitatea curentă a unui club de calculatoare "
  "(sală de gaming / internet café), în care evidența echipamentelor, a membrilor, a sesiunilor "
  "de utilizare și a încasărilor se ține astăzi parțial pe suport de hârtie și în fișiere "
  "Excel, fără o legătură unică între date.")
j("Această situație generează pierderi de timp, dublări de date, greșeli de calcul al costului "
  "sesiunilor, imposibilitatea de a urmări în timp real disponibilitatea echipamentelor și "
  "dificultăți la întocmirea rapoartelor. Pentru a asigura transparența, acuratețea și "
  "eficiența proceselor operaționale, a apărut necesitatea stringentă de automatizare a acestor "
  "activități prin dezvoltarea unei soluții informatice dedicate, instalabile local și "
  "adaptabile fără recompilare.")

h("1.2. Scopul și obiectivele sistemului", 2)
j("Implementarea SIA CCA urmărește atingerea următoarelor obiective strategice și operaționale:")
for i, txt in enumerate([
    "Crearea unui sistem informațional integrat destinat clubului de calculatoare, în vederea "
    "eficientizării proceselor operaționale prin automatizarea activităților de stocare, "
    "sistematizare și actualizare a informațiilor privind sălile, echipamentele, membrii, "
    "sesiunile și rezervările;",
    "Optimizarea activității personalului prin eliminarea proceselor manuale de colectare și "
    "înregistrare a informațiilor pe suport de hârtie și/sau în fișiere Excel, prin "
    "digitalizarea completă a fluxurilor de lucru;",
    "Asigurarea identificării unice și a gestionării centralizate a înregistrărilor, "
    "sesiunilor, rezervărilor și rapoartelor generate în cadrul proceselor operaționale, "
    "garantând trasabilitatea și integritatea informațiilor;",
    "Eliminarea completării manuale a registrelor și rapoartelor prin automatizarea generării "
    "acestora, în baza datelor primare înregistrate în sistem;",
    "Centralizarea datelor din multiple surse (săli, echipamente, membri, operatori) într-o "
    "platformă integrată unică, care permite accesul la informații agregate și detaliate;",
    "Reducerea timpului necesar pentru colectarea, procesarea și consolidarea informațiilor, "
    "precum și pentru generarea rapoartelor statistice și analitice;",
    "Crearea unui spațiu de lucru unificat, care să asigure accesul facil și securizat al "
    "utilizatorilor la resursele și instrumentele sistemului, în funcție de rol.",
], start=1):
    bullet("%d. %s" % (i, txt))

h("1.3. Părțile implicate și rolurile sistemului", 2)

h("1.3.1. Responsabilitatea de interacțiune cu părțile și procesele de afacere asigurate", 3)
j("Partea implicată în asigurarea bunei funcționări a sistemului este administrația clubului de "
  "calculatoare, în calitate de autoritate cointeresată în gestionarea eficientă a sălilor, "
  "echipamentelor, membrilor și încasărilor.")

h("1.3.2. Achizitorul subsistemului informatic", 3)
j("Achizitorul sistemului este clubul de calculatoare (beneficiarul), care efectuează "
  "totalitatea activităților de selectare a executantului, monitorizează procesele de "
  "proiectare/dezvoltare/implementare și recepționează sistemul după faza de acceptanță finală.")

h("1.3.3. Posesorul subsistemului informatic", 3)
j("Posesorul sistemului va fi clubul de calculatoare, care are responsabilitatea de a asigura "
  "gestionarea eficientă a clubului și de a administra datele aplicației.")

h("1.3.4. Administratorul subsistemului informatic", 3)
j("Administratorul sistemului este desemnat de către beneficiar și are scopul administrării "
  "aplicației: gestionarea conturilor și a permisiunilor, configurarea parametrilor, "
  "monitorizarea jurnalului de audit și asigurarea funcționării corecte a bazei de date.")

h("1.3.5. Utilizatorii sistemului", 3)
j("Sistemul prevede mai multe roluri de utilizator, cu drepturi diferențiate. Rolurile de "
  "bază predefinite sunt următoarele:")
table([
    ["Rol", "Descriere și drepturi"],
    ["Administrator", "Acces complet: operații CRUD pe toate entitățile, gestionarea "
     "operatorilor și a utilizatorilor, acordarea/revocarea permisiunilor, rapoarte, "
     "tablou de bord, audit."],
    ["Operator / Manager", "Personalul clubului: înregistrează sesiuni și rezervări, "
     "actualizează inventarul și membrii; rolul de manager are drepturi extinse de vizualizare "
     "și raportare."],
    ["Utilizator (User)", "Acces restrâns: doar vizualizarea echipamentelor, lista de favorite, "
     "profilul propriu și schimbarea parolei."],
    ["Auditor", "Vizualizarea tuturor datelor și a jurnalului de audit, în scop de analiză și "
     "monitorizare, fără drepturi de modificare."],
], widths=[3.8, 11.2])

h("1.4. Terminologii și abrevieri", 2)
j("În cadrul documentului dat au fost utilizate următoarele terminologii:")
for txt in [
    "SIA CCA – Sistemul Informațional Automatizat de gestionare a Clubului de Calculatoare;",
    "CRUD – Create, Read, Update, Delete: operațiile de bază pe date (adăugare, vizualizare, "
    "modificare, ștergere);",
    "Echipament – orice unitate din inventarul clubului: PC gaming, consolă, stație VR, server, "
    "simulator, periferic etc.;",
    "Sesiune – intervalul de utilizare a unui echipament de către un membru, pentru care se "
    "calculează un cost total și o încasare;",
    "Rezervare – programarea anticipată a unui echipament pentru un membru, la o dată și oră "
    "stabilite;",
    "Membru – persoana înregistrată în club (Standard, Premium, Student, Companie);",
    "Audit (jurnal) – înregistrarea automată a operațiilor INSERT/UPDATE/DELETE pentru "
    "trasabilitate;",
    "Permisiune – drept de acces granular acordat unui utilizator pentru o funcție a sistemului;",
]:
    bullet(txt)
j("Prezentul document include descifrarea următoarelor abrevieri:")
for txt in [
    "BD – bază de date;",
    "SGBD – Sistem de Gestiune a Bazelor de Date;",
    "UI – User Interface (interfața cu utilizatorul);",
    "KPI – Key Performance Indicators (indicatori-cheie de performanță);",
    "CSV – Comma-Separated Values (format de fișier pentru export tabelar);",
    "FK – Foreign Key (cheie externă).",
]:
    bullet(txt)

# ==================================================================
#  2. DESCRIEREA SISTEMULUI INFORMAȚIONAL
# ==================================================================
h("2. DESCRIEREA SISTEMULUI INFORMAȚIONAL", 1)
j("SIA CCA va reprezenta o platformă informatică integrată, destinată automatizării proceselor "
  "de evidență, gestionare și monitorizare a activității clubului de calculatoare: săli, "
  "operatori, echipamente, membri, sesiuni de utilizare și rezervări, precum și evidența "
  "încasărilor aferente acestora.")
j("SIA CCA trebuie să fie concepută ca un ecosistem unitar, format dintr-un ansamblu coerent de "
  "componente software și proceduri operaționale, care lucrează împreună pentru a permite "
  "colectarea, procesarea, validarea, stocarea, actualizarea și transmiterea informațiilor "
  "într-un mod automatizat, eficient și sigur, cu intervenție minimă din partea utilizatorilor.")
j("Sistemul este constituit din următoarele componente principale:")
bullet("Proceduri operaționale: fluxurile de lucru adoptate de club (autentificare, sesiune, "
       "rezervare, raportare), în care se integrează interfața software.")
bullet("Software: aplicația desktop AppClubCalculatoare, care asigură interfața de lucru, "
       "serviciile de autentificare, permisiuni, fotografii, favorite, calcul cost, rapoarte și "
       "audit.")
bullet("Bază de date: serverul Microsoft SQL Server cu baza de date relațională "
       "„ClubCalculatoare”, care centralizează toate datele aplicației.")
j("Componentele software solicitate, fluxurile operaționale, rolurile de utilizator și "
  "cerințele tehnice sunt descrise detaliat în capitolele 3–5 ale prezentului document.")

# ==================================================================
#  3. CERINȚELE FAȚĂ DE SISTEM
# ==================================================================
h("3. CERINȚELE FAȚĂ DE SISTEM", 1)
j("Prezentul capitol stabilește cerințele funcționale și nefuncționale obligatorii pe care "
  "soluția informatică trebuie să le îndeplinească pentru a fi acceptată de către beneficiar.")

h("3.1. Cerințe funcționale", 2)
j("Sistemul trebuie să acopere integral procesele de evidență și gestionare a clubului de "
  "calculatoare: înregistrarea, calculul, monitorizarea și raportarea activității. Conform "
  "cerințelor funcționale, procesul de evidență utilizează funcționalitățile de gestionare a "
  "utilizatorilor, gestionare a datelor de domeniu, căutare, raportare, calcul de cost, "
  "jurnalizare și control al accesului, asigurând trasabilitatea completă a datelor.")

h("3.1.1. Cerințe față de arhitectură și componentele sistemului", 3)
j("Aplicația va fi dedicată evidenței și gestionării clubului de calculatoare, incluzând "
  "informații despre săli, echipamente, membri, operatori, sesiuni și rezervări, precum și "
  "aspecte financiare aferente (cost, încasare).")
j("Aplicația va fi dezvoltată ca o soluție de tip desktop, cu interfață grafică nativă "
  "Windows, folosită local în club. Sistemul va comunica direct cu un sistem de gestionare a "
  "bazelor de date relaționale (Microsoft SQL Server), care va stoca informațiile despre "
  "săli, echipamente, membri, operatori, sesiuni, rezervări, utilizatori, fotografii, favorite "
  "și jurnalul de audit.")
j("Aplicația trebuie să fie organizată pe straturi (arhitectură stratificată), astfel:")
bullet("nivel bază de date – tabelele, view-urile, constrângerile de integritate și "
       "interogările;" )
bullet("nivel servicii – componentele care implementează logica: autentificare, permisiuni, "
       "fotografii, favorite, configurare, jurnalizare;")
bullet("nivel prezentare – formularele și controalele interfeței cu utilizatorul.")
j("În cadrul SIA CCA, sistemul trebuie să asigure procesul de evidență cu gestionarea "
  "următoarelor etape principale:")
for i, txt in enumerate([
    "Formarea și actualizarea registrelor de date: sistemul gestionează evidențe distincte "
    "pentru săli, echipamente, membri, operatori, sesiuni și rezervări, în baza datelor "
    "introduse și actualizate de utilizatori;",
    "Determinarea costului/încasării: sistemul calculează costul total al sesiunii și încasarea "
    "aferentă, pe baza tarifului orar și a duratei, cu posibilitatea aplicării de reduceri;",
    "Generarea și vizualizarea informării privind obligațiile: fiecare sesiune și rezervare "
    "conține informații obligatorii (identificator unic, echipamentul, membrul, durata, data, "
    "costul);",
    "Evidența sesiunilor și a încasărilor: utilizatorii înregistrează sesiunile și plățile, "
    "iar sistemul actualizează situația financiară și raportările aferente;",
    "Monitorizarea situațiilor restante și a anomaliilor: sistemul generează automat alerte "
    "privind echipamentele stagnante, rezervările viitoare și sesiunile nesoluționate, precum "
    "și liste de restanțe;"
], start=1):
    bullet("%d. %s" % (i, txt))
j("Sistemul trebuie să ofere utilizatorului un mecanism de autentificare securizat (nume "
  "utilizator și parolă) și să permită accesul diferențiat al utilizatorilor în funcție de "
  "atribuții și responsabilități, pe bază de rol și de permisiuni granulare.")
j("Sistemul trebuie să asigure gestiunea utilizatorilor din cadrul aplicației privind crearea, "
  "modificarea, activarea/dezactivarea conturilor și resetarea parolei.")
j("Sistemul trebuie să dețină un modul flexibil de gestionare a rolurilor și permisiunilor, "
  "care să permită alocarea granulară a drepturilor de vizualizare, modificare, ștergere, "
  "raportare și administrare a datelor din sistem.")
j("Sistemul trebuie să dețină predefinit cel puțin următoarele roluri principale – "
  "Administrator, Operator/Manager, Utilizator și Auditor – fiecare cu permisiuni specifice, "
  "conform proceselor operaționale și responsabilităților organizaționale.")
j("În cadrul SIA CCA, administratorul de sistem trebuie să asigure administrarea tehnică și "
  "operațională a aplicației. Principalele responsabilități ale administratorului sunt:")
for i, txt in enumerate([
    "administrarea parametrilor de sistem și gestionarea conturilor de utilizator, cu "
    "atribuirea drepturilor de acces conform rolurilor;",
    "asigurarea securității sistemului prin implementarea politicilor de parolă, jurnalizare "
    "și monitorizare a accesărilor;",
    "gestionarea copiilor de siguranță (backup) și restaurarea datelor bazei de date;",
    "monitorizarea și auditarea jurnalelor de activitate pentru a preveni, detecta și "
    "documenta acțiunile neautorizate sau neconforme.",
], start=1):
    bullet("%d. %s" % (i, txt))

h("3.1.2. Cerințe față de modulele funcționale", 3)
j("Sistemul informatic trebuie să conțină, obligatoriu, următoarele module funcționale, "
  "descrise mai jos. Cerințele marcate cu prioritate „Obligatoriu” sunt esențiale pentru "
  "acceptanța sistemului.")

p("a) Autentificare și gestionarea accesului", bold=True, space_before=6, space_after=3)
for txt in [
    "Autentificarea utilizatorilor se va efectua obligatoriu înaintea accesării "
    "funcționalităților aplicației, pe bază de nume utilizator și parolă;",
    "Parolele vor fi stocate în baza de date exclusiv sub formă de hash SHA-256 (fără parole în "
    "clar), iar comparația se va face în timp constant (protecție anti atac de tip timing);",
    "Sistemul va limita numărul de tentative eșuate de autentificare și va bloca temporar "
    "contul după depășirea lor, cu afișarea unui indicator vizual;",
    "Sistemul va oferi funcția de resetare a parolei prin e-mail (SMTP) cu generarea unei "
    "parole temporare și funcția de schimbare a parolei pentru orice utilizator autentificat;",
    "Sistemul va închide sesiunea automat la inactivitate, după un interval configurabil.",
]:
    bullet(txt)

p("b) Gestionarea datelor de domeniu (CRUD)", bold=True, space_before=6, space_after=3)
j("Sistemul va oferi formulare generice de adăugare, vizualizare, modificare și ștergere "
  "pentru următoarele tabele de domeniu:")
table([
    ["Entitate", "Câmpuri principale (exemplu)"],
    ["Echipamente", "denumire, tip, sală, putere (W), RAM (GB), an achiziție, status, tarif orar"],
    ["Săli", "denumire, tip (Gaming/VIP/Consolă/VR), zonă/localitate, număr locuri, operator șef"],
    ["Membri", "nume, tip membru (Standard/Premium/Student/Companie), telefon, e-mail"],
    ["Sesiuni", "echipament, membru, operator, data, durata (minute), cost total, încasare"],
    ["Rezervări", "echipament, membru, operator, data programată, status (Programată/...)"],
    ["Operatori", "nume, funcție (operator/tehnician/casier), sală, date de contact"],
], widths=[3.4, 11.6])
for txt in [
    "Fiecare operație de adăugare/modificare/ștergere va fi validată și înregistrată în "
    "jurnalul de audit;",
    "Ștergerea unei înregistrări cu dependențe (chei externe) va fi blocată din aplicație, cu "
    "afișarea unui mesaj explicativ prietenos;",
    "Pentru utilizatorii cu rol restrâns, formularele vor funcționa în mod „doar vizualizare”, "
    "fără butoane de modificare;",
    "Câmpurile opționale vor putea primi valoarea „N/A” (NULL în baza de date).",
]:
    bullet(txt)

p("c) Căutare și filtrare", bold=True, space_before=6, space_after=3)
for txt in [
    "Sistemul va oferi căutare instantanee (filtrare în timp real) în orice listă de date, "
    "după un text introdus de utilizator;",
    "Căutarea va fi aplicată pe toate coloanele text relevante ale entității;",
    "Căutarea nu va modifica datele și nu va permite executarea de cod (protecție anti SQL "
    "injection).",
]:
    bullet(txt)

p("d) Tablou de bord (indicatori KPI) și alerte", bold=True, space_before=6, space_after=3)
for txt in [
    "Sistemul va afișa un tablou de bord cu indicatori-cheie: echipamente libere, sesiuni în "
    "așteptare, rezervări programate, încasare pe perioada selectată, încasare totală și "
    "valoarea portofoliului disponibil;",
    "Tabloul de bord va oferi un selector de perioadă (lună/an), cu construirea automată a "
    "listei de ani din datele existente;",
    "Sistemul va afișa topul operatorilor după încasarea totală și un grafic al încasării pe "
    "șase luni consecutive;",
    "Sistemul va genera alerte operaționale: echipamente libere de peste 90 de zile, rezervări "
    "programate în următoarele 7 zile și sesiuni nesoluționate de peste 30 de zile; pragurile "
    "vor fi configurabile;",
    "Alertele vor fi clicabile și vor deschide lista aferentă, deja filtrată pe înregistrările "
    "semnalate;",
    "Tabloul de bord se va reîmprospăta automat la întoarcerea din orice fereastră de lucru.",
]:
    bullet(txt)

p("e) Rapoarte, statistici, export și tipărire", bold=True, space_before=6, space_after=3)
for txt in [
    "Sistemul va genera rapoarte: echipamente pe săli, sesiuni complete, echipamente "
    "disponibile (cu filtre pe tip și lună, parametrizate), top operatori după încasare și "
    "statistici pe săli;",
    "Sistemul va oferi rapoarte tematice (de exemplu, „PC Gaming” și „Top sesiuni”);",
    "Orice raport afișat va putea fi exportat în format CSV, compatibil cu aplicațiile de tip "
    "foaie de calcul (cu BOM și separator adaptat localizării);",
    "Orice raport va putea fi tipărit sau salvat în PDF, cu previzualizare, antet și "
    "numerotarea paginilor;",
    "Toate interogările din rapoarte vor fi parametrizate (protecție anti SQL injection).",
]:
    bullet(txt)

p("f) Modulul fotografii (opțional)", bold=True, space_before=6, space_after=3)
for txt in [
    "Sistemul va permite asocierea de fotografii la echipamente, cu stocarea fișierelor pe disc "
    "și a metadatelor (căi) în baza de date;",
    "Vor fi generate automat o miniatură și o versiune de dimensiune medie pentru fiecare "
    "imagine;",
    "Vor exista operații de import dintr-un folder, de vizualizare și de ștergere a "
    "fotografiilor, controlate prin permisiuni dedicate.",
]:
    bullet(txt)

p("g) Modulul Favorite (opțional)", bold=True, space_before=6, space_after=3)
for txt in [
    "Utilizatorii vor putea marca echipamentele favorite și le vor regăsi într-o listă "
    "dedicată, cu căutare;",
    "Accesul la modul va fi controlat printr-o permisiune dedicată.",
]:
    bullet(txt)

p("h) Calculatorul de cost (opțional)", bold=True, space_before=6, space_after=3)
for txt in [
    "Sistemul va oferi un calculator pentru estimarea costului unei sesiuni sau al unui "
    "abonament: tarif orar × durată × număr de sesiuni, cu discount de pachet și reducere de "
    "membru;",
    "Rezultatul va include costul brut, reducerea aplicată, totalul de plată și costul mediu "
    "pe sesiune, precum și un tabel comparativ de pachete.",
]:
    bullet(txt)

p("i) Profilul utilizatorului", bold=True, space_before=6, space_after=3)
for txt in [
    "Fiecare utilizator autentificat va putea vizualiza datele contului propriu (nume, e-mail, "
    "data creării, ultima accesare) și lista completă a permisiunilor sale;",
    "Din profil se va putea accesa direct schimbarea parolei.",
]:
    bullet(txt)

p("j) Audit și trasabilitate", bold=True, space_before=6, space_after=3)
for txt in [
    "Fiecare operație de adăugare, modificare și ștergere va fi înregistrată în jurnalul de "
    "audit (tabelă, identificator, acțiune, utilizator, data/ora);",
    "Sistemul va păstra un istoric al acordării/revocării permisiunilor, schimbărilor de rol, "
    "resetărilor de parole și acceselor blocate;",
    "Jurnalele de erori și securitate vor fi scrise într-un fișier de log, cu rotire automată.",
]:
    bullet(txt)

p("k) Interogări obligatorii asupra bazei de date", bold=True, space_before=6, space_after=3)
j("Sistemul, prin baza de date, trebuie să implementeze și să demonstreze cel puțin "
  "următoarele zece interogări/cerințe analitice:")
for i, txt in enumerate([
    "media încasării pe săli;",
    "situația echipamentelor (pe tipuri și status);",
    "membrii care au utilizat toate stațiile de un anumit tip într-o perioadă;",
    "clasamentul sălilor după încasare;",
    "perechi de membri din săli diferite care au folosit același echipament cu același cost;",
    "operatorii cu încasare peste medie;",
    "echipamentele fără rezervări;",
    "statistici lunare ale activității;",
    "cea mai mare creștere a încasării între luni consecutive;",
    "top 3 operatori pe fiecare sală.",
], start=1):
    bullet("%d. %s" % (i, txt))
j("Fiecare dintre aceste interogări trebuie să returneze rezultate pe setul de date "
  "demonstrativ furnizat împreună cu sistemul.")

h("3.1.3. Automatizări și mecanisme de control", 3)
j("Conform cerințelor funcționale, întregul proces trebuie să fie susținut prin:")
for i, txt in enumerate([
    "interfață unificată și intuitivă pentru gestionarea utilizatorilor, datelor și rapoartelor;",
    "clasificatoare interne pentru tipurile de săli, echipamente, membri, sesiuni și rezervări;",
    "exporturi și rapoarte în formate CSV și PDF pentru registre și dări de seamă;",
    "jurnalizarea automată a evenimentelor de business pentru audit și control;",
    "mecanisme de securitate (autentificare, blocarea contului după tentative eșuate, resetare "
    "parolă, expirare sesiune)."
], start=1):
    bullet("%d. %s" % (i, txt))

h("3.2. Cerințe non-funcționale", 2)

h("3.2.1. Cerințe tehnice generale", 3)
for txt in [
    "Interfața aplicației trebuie să fie compatibilă cu sistemele de operare Windows 10 și "
    "Windows 11 (x86/x64) și să suporte afișarea la rezoluții și scalări DPI diferite;",
    "Aplicația trebuie să funcționeze cu o bază de date relațională și să nu depindă de un "
    "browser web;",
    "Toate setările (conexiunea la baza de date, pragurile de securitate și de alertă, căile "
    "fișierelor) trebuie să fie configurabile fără recompilarea aplicației;",
    "Modelul de date al sistemului va fi documentat printr-o schemă (scripturi DDL pentru "
    "SGBD relațional), pentru a asigura trasabilitate și migrare controlată;",
    "Conținutul oricărui tabel de rezultate va putea fi exportat de către utilizatori în CSV "
    "și PDF.",
]:
    bullet(txt)

h("3.2.2. Cerințe față de stiva tehnologică", 3)
j("Pentru asigurarea unei implementări stabile și conforme cu cerințele de performanță și "
  "securitate, sistemul va fi proiectat și implementat în baza următoarelor cerințe:")
table([
    ["Componentă", "Tehnologie solicitată"],
    ["Limbaj de programare", "C# (limbaj orientat pe obiecte)"],
    ["Platformă / runtime", ".NET 8 (Windows Desktop)"],
    ["Interfață grafică", "Windows Forms (WinForms)"],
    ["SGBD", "Microsoft SQL Server 2016 sau mai nou (SQL Server Express acceptat)"],
    ["Acces la date", "Microsoft.Data.SqlClient / System.Data.SqlClient, interogări parametrizate"],
    ["Mediu de dezvoltare", "Visual Studio 2022 sau mai nou, cu workload .NET Desktop"],
    ["Controlul versiunilor", "Git"],
], widths=[4.6, 10.4])
j("Ofertantul va include în propunerea tehnică descrierea stivei tehnologice propuse, "
  "specificând versiunile componentelor software, modulele și librăriile utilizate, "
  "mecanismele de securitate implementate, arhitectura logică a aplicației, precum și soluțiile "
  "de backup și monitorizare.")

h("3.2.3. Cerințe de performanță", 3)
p("Pentru a asigura o funcționare stabilă, performantă și conformă cu nevoile operaționale, "
  "sistemul va respecta următoarele cerințe de performanță:", align="justify")
table([
    ["Indicator", "Țintă"],
    ["Timp de răspuns – operațiuni simple de citire", "maximum 2 secunde"],
    ["Timp de răspuns – operațiuni de scriere/salvare", "maximum 3 secunde"],
    ["Deschiderea listelor de date (până la 1.000 de înregistrări)", "maximum 2 secunde"],
    ["Generarea rapoartelor și export CSV", "maximum 5 secunde"],
    ["Utilizatori simultani (instanțe locale)", "cel puțin 10, fără degradarea performanțelor"],
    ["Volum de date suportat", "cel puțin 100.000 de înregistrări în tabelele de bază"],
], widths=[8.5, 6.5])
j("Pentru operațiunile ce implică volum mare de date, sistemul va utiliza, ori de câte ori "
  "este posibil, mecanisme de procesare asincronă, optimizând timpii de execuție și "
  "experiența utilizatorilor.")

h("3.2.4. Cerințe de integrare", 3)
j("Sistemul va fi proiectat astfel încât să permită integrarea cu infrastructura existentă a "
  "beneficiarului, respectând principiile interoperabilității și reutilizării:")
for txt in [
    "integrare cu serverul Microsoft SQL Server pentru stocarea și regăsirea datelor;",
    "integrare cu un server SMTP pentru funcția de resetare a parolei prin e-mail;",
    "integrabilitate cu aplicațiile de tip foaie de calcul (Excel) prin exportul fișierelor CSV;",
    "reutilizarea identității utilizatorilor din baza de date proprie a sistemului (nume "
    "utilizator și parolă), cu posibilitatea extinderii ulterioare.",
]:
    bullet(txt)

h("3.2.5. Cerințe de securitate și protecție a datelor", 3)
j("Pentru a asigura un nivel înalt de securitate informațională și protecție a datelor, "
  "sistemul va fi proiectat și implementat în conformitate cu cerințele și bunele practici din "
  "domeniul securității cibernetice:")
p("a) Integritatea și confidențialitatea datelor:", bold=True, space_before=4, space_after=3)
for txt in [
    "sistemul trebuie să garanteze păstrarea completă și integritatea datelor pe întreaga "
    "durată a ciclului de viață al acestora;",
    "parolele utilizatorilor vor fi stocate exclusiv sub formă de hash SHA-256, fără parole în "
    "clar;",
    "vor exista proceduri de backup și restaurare a bazei de date.",
]:
    bullet(txt)
p("b) Securitatea aplicației și protecția împotriva vulnerabilităților:", bold=True,
  space_before=4, space_after=3)
for txt in [
    "toate interogările SQL vor fi parametrizate, pentru a preveni atacurile de tip SQL "
    "injection;",
    "toate câmpurile de introducere a datelor vor fi validate, pentru a preveni introducerea "
    "datelor nevalide sau potențial periculoase;",
    "aplicația va fi dezvoltată și testată pentru a fi securizată conform bunelor practici "
    "(eliminarea vulnerabilităților comune: injection, autentificare defectuoasă, expunerea "
    "datelor sensibile, configurare greșită).",
]:
    bullet(txt)
p("c) Controlul accesului și autentificarea utilizatorilor:", bold=True, space_before=4,
  space_after=3)
for txt in [
    "accesul la sistem va fi realizat în mod controlat și securizat, pe bază de rol și "
    "permisiuni;",
    "autentificarea va fi obligatorie înaintea accesării funcționalităților, iar drepturile de "
    "acces vor fi acordate conform principiilor „need to know” și „least privilege”;",
    "sistemul va implementa blocarea automată a contului după un număr configurabil de "
    "tentative eșuate și expirarea sesiunii la inactivitate.",
]:
    bullet(txt)
p("d) Jurnalizare și trasabilitate:", bold=True, space_before=4, space_after=3)
for txt in [
    "toate acțiunile de modificare a datelor vor fi înregistrate în jurnalul de audit "
    "(utilizatorul, acțiunea, data și ora, tabela, identificatorul);",
    "jurnalele de erori și de securitate vor fi păstrate într-un fișier de log cu rotire "
    "automată;",
    "jurnalele vor permite auditul complet al activităților.",
]:
    bullet(txt)

h("3.2.6. Cerințe de instruire", 3)
for txt in [
    "Executantul va organiza sesiuni de instruire dedicate, atât teoretice, cât și practice, "
    "pentru fiecare grup de utilizatori (operatori, administratori, personal de suport);",
    "Materialele de instruire (ghiduri, fișe de lucru, manuale) vor fi elaborate în limba "
    "română și vor reflecta terminologia, structura și funcționalitățile sistemului;",
    "Materialele vor fi livrate beneficiarului în format electronic editabil și/sau tipărit.",
]:
    bullet(txt)

h("3.2.7. Cerințe de garanție", 3)
j("Pentru a asigura funcționarea stabilă, corectă și continuă a soluției livrate, executantul "
  "va acorda o perioadă de garanție de 3 luni, care începe imediat după punerea în funcțiune "
  "finală a sistemului. Perioada de pilotare destinată testării și validării nu se include în "
  "perioada de garanție.")
p("a) Obligațiile executantului în perioada de garanție:", bold=True, space_before=4,
  space_after=3)
j("Pe întreaga durată a perioadei de garanție, executantul va asigura suport tehnic continuu, "
  "eliminarea defecțiunilor raportate și rezolvarea incidentelor în conformitate cu un acord de "
  "nivel de serviciu (SLA). Timpul maxim de răspuns și intervenție va respecta următoarele "
  "praguri:")
bullet("pentru erorile critice: răspuns în maximum 4 ore, soluționare în termen de 1 zi "
       "lucrătoare;")
bullet("pentru erorile non-critice: răspuns în maximum 24 de ore, soluționare în termen de 5 "
       "zile lucrătoare.")
p("b) Trasabilitatea și documentarea incidentelor:", bold=True, space_before=4, space_after=3)
j("Executantul va implementa un sistem de documentare și monitorizare a incidentelor, care va "
  "asigura trasabilitatea completă a fiecărei sesizări, de la înregistrare până la soluționare, "
  "inclusiv urmărirea statusului și arhivarea soluțiilor aplicate.")

h("3.2.8. Cerința față de documentație", 3)
j("Pentru a asigura exploatarea corectă, sigură și eficientă a sistemului, executantul are "
  "obligația de a elabora, structura și livra un set complet de documentații tehnice și "
  "operaționale, atât pentru utilizatorii finali, cât și pentru echipa tehnică:")
for txt in [
    "Ghiduri și manuale de utilizator, adaptate fiecărui rol, conținând instrucțiuni pas cu pas, "
    "capturi de ecran și exemple de utilizare;",
    "Documentația tehnică: sarcina tehnică, prototipul sistemului, planul de testare, procesele-"
    "verbale de testare;",
    "Manualul de instalare și configurare, care va include: instrucțiunile de compilare a "
    "codului sursă, cerințele hardware și software minime și recomandate, descrierea și "
    "configurarea bazei de date, procedurile de backup și recuperare.",
]:
    bullet(txt)
j("Documentația tehnică va fi livrată în limba română, în format electronic (PDF editabil sau "
  "DOCX), asigurând lizibilitatea și utilizabilitatea acesteia.")

h("3.2.9. Cerințe pentru testarea de acceptanță a sistemului", 3)
j("Pentru a garanta calitatea, stabilitatea și conformitatea sistemului cu cerințele "
  "specificate, executantul are obligația de a organiza și desfășura un proces complet de "
  "testare și acceptanță, în colaborare cu beneficiarul.")
p("a) Organizarea procesului de testare:", bold=True, space_before=4, space_after=3)
j("Executantul va fi responsabil de organizarea testării de acceptanță, prin:")
for txt in [
    "elaborarea planului detaliat de testare (obiective, metodologie, mediu, criterii de "
    "acceptanță, calendar);",
    "pregătirea scenariilor de testare, bazate pe cazurile de utilizare și fluxurile "
    "operaționale;",
    "executarea testelor și documentarea erorilor identificate, cu clasificarea pe severitate "
    "(critice, majore, minore);",
    "corectarea erorilor și reluarea testelor aferente, până la obținerea conformității "
    "complete;",
    "elaborarea raportului final de testare, cu rezultatele obținute și concluzia privind "
    "nivelul de conformitate.",
]:
    bullet(txt)
p("b) Tipuri de testare obligatorii:", bold=True, space_before=4, space_after=3)
for txt in [
    "Testarea funcțională – verificarea faptului că toate funcționalitățile se comportă conform "
    "cerințelor din caietul de sarcini;",
    "Testarea de integrare – verificarea interacțiunii dintre aplicație și baza de date;",
    "Testarea de validare – verificarea validării datelor introduse (cazuri pozitive și "
    "negative);",
    "Testarea de acces pe roluri – verificarea restricțiilor de acces pentru fiecare rol;",
    "Testarea de securitate – verificarea protecției împotriva atacurilor de tip SQL injection "
    "și a autentificării defectuoase.",
]:
    bullet(txt)
p("c) Criterii de acceptanță:", bold=True, space_before=4, space_after=3)
j("Sistemul va fi considerat acceptat doar în cazul în care:")
for txt in [
    "nu există erori critice (blocante) rămase nerezolvate;",
    "numărul erorilor majore identificate este mai mic de 3;",
    "toate erorile minore sunt documentate și incluse în planul de remediere.",
]:
    bullet(txt)
j("După îndeplinirea acestor condiții, se va proceda la semnarea procesului-verbal de "
  "acceptanță.")

h("3.2.10. Cerințele față de livrabilele sistemului", 3)
j("Livrabilele proiectului reprezintă totalitatea artefactelor software, documentațiilor "
  "tehnice și operaționale, precum și materialele aferente transferului de cunoștințe către "
  "beneficiar. Toate livrabilele vor fi furnizate în limba română, în formate deschise și "
  "accesibile.")
p("a) Componența livrabilelor:", bold=True, space_before=4, space_after=3)
table([
    ["Nr.", "Livrabil", "Descriere"],
    ["1", "Codul sursă complet", "Soluția compilabilă, cu proiectul Visual Studio, comentat și "
     "documentat."],
    ["2", "Scripturile bazei de date", "Scripturi DDL/DML pentru crearea bazei de date, a "
     "tabelelor, view-urilor și încărcarea datelor demonstrative."],
    ["3", "Sarcina tehnică", "Documentul care descrie cerințele, arhitectura și modulele "
     "sistemului."],
    ["4", "Prototipul sistemului", "Structura interfeței grafice și fluxurile principale."],
    ["5", "Planul de testare", "Strategia, scenariile și criteriile de acceptanță."],
    ["6", "Procese-verbale de testare", "Documente ce confirmă executarea și rezultatele "
     "testelor."],
    ["7", "Manualul de instalare și configurare", "Instrucțiuni de instalare, compilare și "
     "configurare a bazei de date."],
    ["8", "Manualele utilizatorului", "Ghiduri pe roluri, cu capturi de ecran, în format PDF."],
], widths=[1.2, 4.6, 9.2])
p("b) Procedura de examinare și acceptanță a livrabilelor:", bold=True, space_before=4,
  space_after=3)
j("Pentru fiecare livrabil, executantul va respecta o procedură clară de predare și validare:")
bullet("transmiterea livrabilului spre examinare cu cel puțin 5 zile lucrătoare înainte de "
       "data preconizată a acceptanței;")
bullet("perioada de examinare, în care executantul furnizează clarificări la solicitare;")
bullet("feedback și ajustări – beneficiarul transmite comentarii explicite privind "
       "neconformitățile constatate;")
bullet("remediere – executantul soluționează deficiențele și retransmite livrabilul revizuit în "
       "maximum 5 zile lucrătoare;")
bullet("acceptanța finală – livrabilul este considerat acceptat la semnarea documentului de "
       "acceptanță.")

# ==================================================================
#  4. CERINȚE FAȚĂ DE OFERTANȚI
# ==================================================================
h("4. CERINȚE FAȚĂ DE OFERTANȚI", 1)
j("În scopul demonstrării capacității profesionale, tehnice și a experienței necesare pentru "
  "executarea corespunzătoare a contractului, ofertanții trebuie să îndeplinească cerințele de "
  "calificare și eligibilitate stabilite în prezenta secțiune.")

h("4.1. Cerințe față de competențele tehnice ale ofertantului", 2)
j("Ofertantul trebuie să demonstreze experiență relevantă în elaborarea aplicațiilor desktop și "
  "a sistemelor informaționale conectate la baze de date relaționale. Totodată, acesta va "
  "prezenta dovezi privind experiența acumulată, prin indicarea proiectelor similare livrate "
  "anterior, care să reflecte capacitatea tehnică de realizare a serviciilor solicitate.")

h("4.2. Cerințe față de echipa de implementare", 2)
j("Ofertantul trebuie să demonstreze că dispune de o echipă de implementare profesionistă, cu "
  "experiență relevantă pentru prestarea serviciilor solicitate. Echipa de implementare va "
  "include cel puțin următoarele funcții și competențe:")
table([
    ["Funcție", "Persoană", "Competențe solicitate"],
    ["Manager de proiect", "Dudnic Alexandru", "Coordonarea activităților, planificarea "
     "sarcinilor, comunicarea cu beneficiarul, respectarea termenelor."],
    ["Analist / Proiectant", "Marinescu Mihail", "Analiza cerințelor, proiectarea structurii "
     "bazei de date și a modulelor funcționale, prototiparea interfeței."],
    ["Programator software", "Scripnic Vladlen", "Dezvoltarea aplicației în C#/WinForms, "
     "implementarea serviciilor, integrarea cu baza de date SQL Server."],
    ["Programator / Testare", "Dudnic Alexandru", "Implementarea modulelor de raportare și "
     "audit, executarea testelor, documentarea rezultatelor."],
], widths=[3.4, 3.4, 8.2])
j("Pentru fiecare membru al echipei, ofertantul va prezenta un Curriculum Vitae (CV) "
  "actualizat, care să reflecte experiența, competențele și certificările deținute.")

h("4.3. Cerințe față de managementul proiectului", 2)
j("Ofertantul trebuie să asigure implementarea proiectului în baza unei metode moderne de "
  "dezvoltare (Agile), utilizând metodologii consacrate precum KANBAN și SCRUM, care să permită "
  "o planificare iterativă, o colaborare activă cu beneficiarul și o adaptare flexibilă la "
  "cerințele funcționale.")
j("Sistemul informatic urmează a fi proiectat și dezvoltat în baza unei abordări hibride de "
  "gestiune a proiectului, care combină elemente de planificare predictivă (Waterfall) cu "
  "procese iterative și incrementale specifice metodologiei Agile.")
j("Iterațiile (sprinturile) de dezvoltare vor avea o durată standard de două săptămâni. La "
  "finalul fiecărei iterații, ofertantul va elabora rapoarte periodice de progres, care vor "
  "include cel puțin:")
bullet("lista funcționalităților implementate în cadrul sprintului curent;")
bullet("lista funcționalităților planificate pentru sprintul următor;")
bullet("prezentarea problemelor apărute, împreună cu măsurile propuse pentru soluționare.")
j("Atât beneficiarul, cât și ofertantul, vor desemna câte un responsabil de proiect, care va "
  "avea rolul de a coordona activitățile și de a asigura o comunicare permanentă între părți, "
  "în vederea bunei desfășurări a lucrărilor și respectării termenelor asumate.")

# ==================================================================
#  5. CERINȚE FAȚĂ DE OFERTE
# ==================================================================
h("5. CERINȚE FAȚĂ DE OFERTE", 1)

h("5.1. Cerințe față de propunerea tehnică", 2)
j("Propunerea tehnică trebuie să fie elaborată și prezentată integral în limba română, "
  "redactată într-o manieră clară, structurată și completă, astfel încât să demonstreze "
  "înțelegerea obiectivelor proiectului și conformitatea cu cerințele din caietul de sarcini.")
j("Ofertanții vor descrie în detaliu arhitectura sistemului propus, prin prezentarea "
  "diagramelor funcționale corespunzătoare fiecărui proces descris ca modul funcțional. "
  "Propunerea tehnică va include o descriere clară a componentelor sistemului și modul în care "
  "acestea răspund cerințelor din caietul de sarcini.")
j("Pentru fiecare caz de utilizare identificat, ofertanții vor prezenta diagrame de proces "
  "(UML / BPMN) care să reflecte logica și fluxurile operaționale propuse. În plus, se solicită "
  "includerea unei viziuni grafice asupra interfeței sistemului, prin prezentarea a minimum "
  "3–5 ecrane exemplificative ale principalelor module funcționale (a se vedea Anexa 2).")
j("Propunerea tehnică va conține, de asemenea, un plan detaliat de implementare a sistemului, "
  "care va include lista activităților, a livrabilelor, termenelor de execuție și "
  "responsabililor desemnați pentru fiecare etapă. Se vor accepta doar propunerile tehnice "
  "complete, care răspund integral cerințelor din prezentul caiet de sarcini.")
j("Durata de realizare a proiectului este de până la 12 săptămâni din data semnării "
  "contractului, perioadă în care ofertantul va asigura finalizarea tuturor activităților "
  "planificate și livrarea completă a sistemului informatic.")

h("5.2. Cerință față de oferta financiară", 2)
j("Oferta financiară trebuie să reflecte integral costurile aferente tuturor serviciilor și "
  "livrabilelor prevăzute în caietul de sarcini. Suma totală ofertată va include toate "
  "cheltuielile necesare pentru executarea completă a contractului, inclusiv costurile de "
  "analiză, proiectare, dezvoltare, testare, implementare, instruire, documentare și "
  "mentenanță în perioada de garanție.")
j("Oferta financiară va fi detaliată pe etape de implementare, prezentând pentru fiecare etapă: "
  "denumirea activităților, livrabilele corespunzătoare, termenele de realizare, valoarea "
  "parțială și procentul aferent din valoarea totală a contractului. Această structură va "
  "permite corelarea clară între progresul tehnic și cel financiar.")
table([
    ["Etapă", "Activități principale", "Livrabil", "Pondere (%)"],
    ["1. Analiză și proiectare", "Analiza cerințelor, proiectarea bazei de date și a "
     "modulelor, prototipul interfeței", "Sarcina tehnică, Prototipul", "20%"],
    ["2. Dezvoltare", "Implementarea modulelor funcționale, a serviciilor și a bazei de date",
     "Cod sursă, Scripturi BD", "40%"],
    ["3. Testare", "Elaborarea și executarea testelor, remedierea defectelor",
     "Plan de testare, Procese-verbale", "20%"],
    ["4. Implementare și instruire", "Instalarea, configurarea, instruirea utilizatorilor, "
     "documentația", "Manual de instalare, Manuale de utilizator", "20%"],
], widths=[3.4, 5.4, 4.6, 1.6])
j("Plățile vor fi efectuate de către beneficiar proporțional cu etapele finalizate și "
  "livrabilele recepționate, în baza proceselor-verbale de recepție semnate de ambele părți, "
  "conform graficului financiar asumat prin contract.")

# ==================================================================
#  ANEXA 1 — STRUCTURA BAZEI DE DATE
# ==================================================================
doc.add_page_break()
h("ANEXA 1. Structura bazei de date „ClubCalculatoare”", 1)
j("Baza de date relațională conține tabelele de domeniu și tabelele de suport (utilizatori, "
  "permisiuni, audit, favorite, fotografii), împreună cu view-urile și interogările analitice:")
table([
    ["Tabelă / View", "Rol"],
    ["Operatori", "personalul clubului (operatori, tehnicieni, casieri)"],
    ["Sali", "sălile clubului (Gaming, VIP, Consolă, VR)"],
    ["Echipamente", "inventarul (PC gaming, console, stații VR, servere, simulatoare)"],
    ["Membri", "membrii clubului (Standard, Premium, Student, Companie)"],
    ["Sesiuni", "sesiunile de utilizare a echipamentelor, cu durată, cost și încasare"],
    ["Rezervari", "rezervările de echipamente, cu data programată și status"],
    ["Utilizatori", "conturile de acces (nume, parolă hash, rol, e-mail, activ)"],
    ["Permisiuni / UtilizatoriPermisiuni", "catalogul drepturilor de acces și acordările per "
     "utilizator"],
    ["Favorite", "echipamentele favorite ale fiecărui utilizator"],
    ["EchipamentFoto", "metadatele fotografiilor echipamentelor (fișierele stau pe disc)"],
    ["AuditLog", "jurnalul operațiilor INSERT/UPDATE/DELETE din aplicație"],
    ["vw_EchipamenteComplete", "echipament + sală + operatorul șef al sălii"],
    ["vw_SesiuniComplete", "sesiune + echipament + membru + operator + sală"],
    ["vw_EchipamenteDisponibile", "doar echipamentele cu status „Liber”"],
    ["vw_MembriCuSesiuni", "membri + numărul și valoarea totală a sesiunilor"],
], widths=[5.4, 9.6])
j("Toate cheile externe sunt definite cu regula „ON DELETE NO ACTION”, iar ștergerea unei "
  "înregistrări cu dependențe este blocată din aplicație, cu afișarea unui mesaj explicativ.")

# ==================================================================
#  ANEXA 2 — INTERFAȚA GRAFICĂ (PROTOTIP)
# ==================================================================
doc.add_page_break()
h("ANEXA 2. Interfața grafică (prototip / ecrane exemplificative)", 1)
j("Mai jos sunt prezentate ecrane exemplificative ale principalelor module funcționale ale "
  "sistemului, care ilustrează aspectul și structura interfeței.")
figure("01_login_admin_dashboard.png",
       "Fig. 1. Autentificare și tabloul de bord al administratorului cu indicatori KPI.")
figure("02_echipamente_lista.png",
       "Fig. 2. Lista de echipamente, cu căutare instantanee și operații CRUD.")
figure("15_sali_lista.png",
       "Fig. 3. Gestionarea sălilor clubului.")
figure("20_rapoarte.png",
       "Fig. 4. Modulul de rapoarte și statistici, cu export CSV și tipărire.")
figure("25_utilizatori.png",
       "Fig. 5. Gestionarea utilizatorilor și a permisiunilor.")
figure("28_calculator.png",
       "Fig. 6. Calculatorul de cost al sesiunii/abonamentului.")

# ---------- finalizare ----------
add_footer_page()
enable_update_fields()
doc.save(OUT)
print("SAVED:", OUT)
print("PARAGRAFE:", len(doc.paragraphs))